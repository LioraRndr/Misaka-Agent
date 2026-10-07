"""Pictures in a deck, a Word document and an EPUB: named where the model reads, shown on request.

2026-10-07: a slide whose meaning is in a picture carried ``needs_vlm: true`` and nothing could
show it; a Word document's pictures were an ``![alt]`` marker with nothing behind it; an EPUB's
were dropped from the text outright. A deck now goes the way FrontierAgent sends one to its
vision reader -- exported to PDF by LibreOffice, one page rendered -- and a picture a Word
document or an EPUB embeds is shown as the file it is stored as."""
import base64
import io
import zipfile
from types import SimpleNamespace as NS

import pytest
from PIL import Image

from misaka.core.documents import index as corpus
from misaka.core.documents.wiring import documents
from misaka.core.wiring import ToolCollector


def _png(width, height, shade=90):
    buffer = io.BytesIO()
    Image.new("L", (width, height), shade).save(buffer, format="PNG")
    return buffer.getvalue()


async def _call(tool_name, args, workspace):
    collector = ToolCollector()
    documents.register(collector)
    tool = next(t for t in collector.tools if t.name == tool_name)
    return await tool.execute("call", args, None, None, NS(cwd=workspace, model=None, modelRegistry=None, signal=None))


def _ingest(path, workspace):
    doc_id, _ = corpus.ingest(str(path), workspace=str(workspace), with_tree=False)
    return doc_id


def _shown(result):
    image = next(part for part in result["content"] if getattr(part, "mimeType", None))
    return Image.open(io.BytesIO(base64.b64decode(image.data)))


# -- where a picture's marker puts it ------------------------------------------------------------------

def test_pictures_are_put_on_the_page_their_marker_is_on():
    pages = ["Intro ![Map of Russia] text", "no pictures here", "![image 2] and ![Plate 3]"]
    listed = [{"marker": "![Map of Russia]", "alt": "Map of Russia", "member": "a.png"},
              {"marker": "![image 2]", "alt": "", "member": "b.png"},
              {"marker": "![Plate 3]", "alt": "Plate 3", "member": None}]
    found = corpus._embedded_inventory(pages, listed)
    assert {page: [i["member"] for i in r["embedded"]] for page, r in found.items()} == {1: ["a.png"], 3: ["b.png", None]}


def test_a_marker_the_pages_lost_stops_the_list_rather_than_misplacing_it():
    listed = [{"marker": "![a]", "alt": "a", "member": "a.png"}, {"marker": "![gone]", "alt": "", "member": "g.png"},
              {"marker": "![b]", "alt": "b", "member": "b.png"}]
    assert list(corpus._embedded_inventory(["![a]", "![b]"], listed)) == [1]


# -- a Word document ----------------------------------------------------------------------------------------

async def test_a_word_document_shows_each_picture_it_embeds(tmp_path):
    import docx
    workspace = tmp_path / "ws"
    workspace.mkdir()
    (tmp_path / "wide.png").write_bytes(_png(300, 120))
    (tmp_path / "tall.png").write_bytes(_png(90, 200, shade=40))
    document = docx.Document()
    document.add_paragraph("The march began at the Niemen in June 1812.")
    document.add_picture(str(tmp_path / "wide.png"))
    document.add_paragraph("The second picture shows the bridge at the Berezina.")
    document.add_picture(str(tmp_path / "tall.png"))
    document.save(str(workspace / "march.docx"))
    doc_id = _ingest(workspace / "march.docx", workspace)
    found = corpus.figures(doc_id, workspace=str(workspace))
    assert [len(r["embedded"]) for r in found.values()] == [2]
    page = next(iter(found))
    notes = (await _call("doc_read", {"doc_id": doc_id, "pages": str(page)}, str(workspace)))["content"][0]["text"]
    assert f"Page {page} holds 2 pictures the document embeds" in notes
    second = await _call("doc_page_image", {"doc_id": doc_id, "page": page, "figure": 2}, str(workspace))
    assert _shown(second).size == (90, 200)
    with pytest.raises(ValueError, match="no printed page"):
        await _call("doc_page_image", {"doc_id": doc_id, "page": page}, str(workspace))


# -- an EPUB ------------------------------------------------------------------------------------------------

CHAPTER = ('<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapter One</h1>'
           '<p>The army crossed the river at dawn, and the map below shows where.</p>'
           '<p><img src="../Images/map%201.png" alt="Map of the crossing"/></p></body></html>')


def _epub(path, chapter=CHAPTER):
    with zipfile.ZipFile(path, "w") as book:
        book.writestr("mimetype", "application/epub+zip")
        book.writestr("META-INF/container.xml",
                      '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                      '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        book.writestr("OEBPS/content.opf",
                      '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                      '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>The March</dc:title></metadata>'
                      '<manifest><item id="c1" href="Text/ch1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                      '<spine><itemref idref="c1"/></spine></package>')
        book.writestr("OEBPS/Text/ch1.xhtml", chapter)
        book.writestr("OEBPS/Images/map 1.png", _png(240, 160))
    return path


async def test_an_epub_marks_its_pictures_and_shows_them(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    doc_id = _ingest(_epub(workspace / "march.epub"), workspace)
    assert "![Map of the crossing]" in corpus.read_page(doc_id, 1, workspace=str(workspace))
    assert corpus.figures(doc_id, workspace=str(workspace))[1]["embedded"][0]["member"] == "OEBPS/Images/map 1.png"
    shown = await _call("doc_page_image", {"doc_id": doc_id, "page": 1, "figure": 1}, str(workspace))
    assert _shown(shown).size == (240, 160)


def test_a_web_page_is_not_given_picture_markers():
    from misaka.core.documents import htmltext
    assert "![" not in htmltext.readable('<p>Price <img src="/icon.png" alt="cart"> 12 EUR</p>')[0]


# -- a deck -------------------------------------------------------------------------------------------------

def _deck(path, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    picture = tmp_path / "chart.png"
    picture.write_bytes(_png(400, 300))
    deck = Presentation()
    for n, kind in enumerate(("text", "picture", "hidden", "picture"), 1):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Slide title {n}"
        if kind == "picture":
            slide.shapes.add_picture(str(picture), Inches(1), Inches(2))
        if kind == "hidden":
            slide._element.set("show", "0")
    deck.save(str(path))
    return path


def _page_text(pdf_path, page):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_path)
    try:
        return pdf[page - 1].get_textpage().get_text_bounded()
    finally:
        pdf.close()


def test_a_slide_that_needs_a_vision_model_is_named(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    doc_id = _ingest(_deck(workspace / "talk.pptx", tmp_path), workspace)
    assert sorted(r["slide"] for r in corpus.figures(doc_id, workspace=str(workspace)).values()) == [2, 4]


async def test_a_slide_is_rendered_through_libreoffice_past_the_hidden_one(tmp_path):
    from misaka.core.documents.office import soffice
    if soffice.binary() is None:
        pytest.skip("LibreOffice is not installed")
    workspace = tmp_path / "ws"
    workspace.mkdir()
    doc_id = _ingest(_deck(workspace / "talk.pptx", tmp_path), workspace)
    page = next(p for p, r in corpus.figures(doc_id, workspace=str(workspace)).items() if r["slide"] == 4)
    result = await _call("doc_page_image", {"doc_id": doc_id, "page": page}, str(workspace))
    assert "(slide 4)" in result["content"][0]["text"] and _shown(result).size[0] > 0
    rendered = corpus.resolve_doc(doc_id, workspace=str(workspace)) + "/rendered.pdf"
    assert "Slide title 4" in _page_text(rendered, 3), "the export's third page is slide 4: hidden slide 3 is not in it"
    hidden = next(p for p in range(1, 10) if "## Slide 3" in (corpus.read_page(doc_id, p, workspace=str(workspace)) or ""))
    with pytest.raises(ValueError, match="hidden"):
        await _call("doc_page_image", {"doc_id": doc_id, "page": hidden}, str(workspace))


def test_an_epub_indexed_before_markers_is_read_again_and_its_pictures_named(tmp_path):
    import json
    workspace = tmp_path / "ws"
    workspace.mkdir()
    book = _epub(workspace / "march.epub")
    doc_id = _ingest(book, workspace)
    ddir = corpus.resolve_doc(doc_id, workspace=str(workspace))
    page = f"{ddir}/pages/p0001.txt"
    with open(page, encoding="utf-8") as f:
        text = f.read()
    with open(page, "w", encoding="utf-8") as f:
        f.write(text.replace("![Map of the crossing]", ""))            # what version 1 stored
    with open(f"{ddir}/meta.json", encoding="utf-8") as f:
        meta = json.load(f)
    meta.pop("extract_version")
    with open(f"{ddir}/meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f)
    assert corpus.figures(doc_id, workspace=str(workspace)) == {}      # and found nothing to name
    _ingest(book, workspace)
    assert "![Map of the crossing]" in corpus.read_page(doc_id, 1, workspace=str(workspace))
    assert corpus.figures(doc_id, workspace=str(workspace))[1]["embedded"][0]["member"] == "OEBPS/Images/map 1.png"
    with open(f"{ddir}/meta.json", encoding="utf-8") as f:
        assert json.load(f)["reread"]["pages"] == [1]


def _as_version_2(ddir, marker):
    """What version 2 stored: the page without ``marker``, stamped as extracted by version 2."""
    import json
    page = f"{ddir}/pages/p0001.txt"
    with open(page, encoding="utf-8") as f:
        text = f.read()
    with open(page, "w", encoding="utf-8") as f:
        f.write(text.replace(marker, ""))
    with open(f"{ddir}/meta.json", encoding="utf-8") as f:
        meta = json.load(f)
    meta["extract_version"] = 2
    with open(f"{ddir}/meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f)


async def test_a_cover_framed_in_svg_is_marked_and_an_epub_indexed_without_it_is_read_again(tmp_path):
    """Calibre writes a cover or a plate as <svg><image xlink:href=...>; the svg was hidden as a
    drawing, and the picture with it."""
    workspace = tmp_path / "ws"
    workspace.mkdir()
    book = _epub(workspace / "march.epub",
                 '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:xlink="http://www.w3.org/1999/xlink">'
                 '<body><h1>Chapter One</h1><svg viewBox="0 0 240 160"><image width="240" height="160" '
                 'xlink:href="../Images/map%201.png"/></svg><p>The army crossed the river at dawn.</p></body></html>')
    doc_id = _ingest(book, workspace)
    assert "![image 1]" in corpus.read_page(doc_id, 1, workspace=str(workspace))
    shown = await _call("doc_page_image", {"doc_id": doc_id, "page": 1, "figure": 1}, str(workspace))
    assert _shown(shown).size == (240, 160)
    _as_version_2(corpus.resolve_doc(doc_id, workspace=str(workspace)), "![image 1]")
    _ingest(book, workspace)
    assert "![image 1]" in corpus.read_page(doc_id, 1, workspace=str(workspace))


def test_a_picture_too_large_to_decode_here_is_refused_in_a_sentence():
    """41 KB of PNG, 12000 x 12000 pixels: 800 MB in the agent's process."""
    from misaka.core.documents.wiring import documents as tools
    huge = io.BytesIO()
    Image.new("1", (12000, 12000)).save(huge, format="PNG")
    with pytest.raises(ValueError, match="12000 x 12000 pixels, too large"):
        tools._pillow_png(huge.getvalue())
    photo = io.BytesIO()
    Image.new("RGB", (9000, 6000), "white").save(photo, format="JPEG")
    shown = Image.open(io.BytesIO(tools._pillow_png(photo.getvalue())))
    assert shown.size[0] * shown.size[1] <= tools.PICTURE_PIXELS_AT_MOST, "a JPEG is decoded smaller"
